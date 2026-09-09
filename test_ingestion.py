import docx
import os
import sys

from rag_core.ingestion import ResumeExtractor

def create_dummy_resume(filename="dummy_resume.docx"):
    doc = docx.Document()
    doc.add_heading("John Doe", 0)
    doc.add_heading("Summary", level=1)
    doc.add_paragraph("Experienced software engineer with 5 years in Python and Machine Learning.")
    
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Senior AI Engineer - Tech Corp (2020 - Present)\n- Built RAG chatbot using LangChain and HuggingFace.\n- Optimized data pipelines.")
    
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.S. Computer Science - State University (2015-2019)")
    
    doc.save(filename)
    print(f"Created {filename}")

if __name__ == "__main__":
    create_dummy_resume()
    extractor = ResumeExtractor()
    text = extractor.extract_text("dummy_resume.docx")
    print("--- Extracted Text ---")
    print(text)
    print("----------------------")
    if text and "John Doe" in text:
        print("Test passed!")
        sys.exit(0)
    else:
        print("Test failed!")
        sys.exit(1)
